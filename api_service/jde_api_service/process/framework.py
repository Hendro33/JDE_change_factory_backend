"""
Process frameworks: a company's imported process hierarchy.

A framework is imported from a workbook (.xlsx): either content the
customer is authorised to use (an APQC Process Classification Framework
edition they licensed) or a hierarchy the customer defines. Jade never
supplies APQC content itself and never invents official identifiers --
external references are shown exactly as the workbook supplied them.

Import is upload -> column mapping -> validation and preview (a draft) ->
activation by an administrator. An activated version is immutable; nodes
of every version are kept, so a story mapped to version 1 still resolves
exactly after version 2 is activated. The original file is stored with
its checksum as provenance.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
from datetime import datetime, timezone
from typing import Any, Optional

from ..discovery import artifacts
from ..persistence.db import connection

SOURCE_KINDS = {
    "apqc_authorised": "APQC content the customer is authorised to use",
    "customer_defined": "Customer-defined hierarchy",
    "synthetic_fixture": "SYNTHETIC fixture for demonstration -- not APQC content",
}
FIELDS = ("node_key", "parent_key", "level", "name", "description", "node_type", "external_ref")
REQUIRED_FIELDS = ("node_key", "name")
FIELD_LABELS = {
    "node_key": "Node ID", "parent_key": "Parent ID", "level": "Level", "name": "Name",
    "description": "Description", "node_type": "Node type", "external_ref": "External reference",
}
# Header spellings recognised when proposing a column mapping.
_ALIASES = {
    "node_key": ("node id", "id", "hierarchy id", "process id", "code", "number"),
    "parent_key": ("parent id", "parent", "parent code"),
    "level": ("level", "depth"),
    "name": ("name", "process name", "element", "title"),
    "description": ("description", "definition"),
    "node_type": ("node type", "type", "category"),
    "external_ref": ("external reference", "pcf id", "external id", "reference"),
}
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20000
TEMPLATE_SHEET = "Framework"


class FrameworkRefused(Exception):
    pass


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def node_sha(node: dict) -> str:
    material = {k: node.get(k) or "" for k in ("node_key", "parent_key", "name", "description", "node_type",
                                              "external_ref")}
    return hashlib.sha256(json.dumps(material, sort_keys=True).encode()).hexdigest()


# ---------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------
def template_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = TEMPLATE_SHEET
    ws.append([FIELD_LABELS[f] for f in FIELDS])
    ws.append(["1", "", 1, "Example category", "Replace these example rows with your framework", "Category", ""])
    ws.append(["1.1", "1", 2, "Example process group", "", "Process group", ""])
    ws.append(["1.1.1", "1.1", 3, "Example process", "", "Process", ""])
    notes = wb.create_sheet("Instructions")
    for line in (
        "Jade process framework import template.",
        "One row per node. Node ID and Name are required; Node IDs must be unique.",
        "Parent ID links a node to its parent. If you leave it empty, Jade can derive the parent from a dotted "
        "Node ID (1.2.3 -> 1.2) when you choose that option during import.",
        "External reference: an identifier from the source framework (for example an APQC PCF ID), only if "
        "your organisation is authorised to use that content. Jade shows it exactly as supplied and does not "
        "verify it.",
        "You can use your own column headings: you map them to these fields during import.",
    ):
        notes.append([line])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------
# Reading a workbook
# ---------------------------------------------------------------------
def _read_workbook(data: bytes) -> dict[str, list[list[Any]]]:
    from openpyxl import load_workbook

    if len(data) > MAX_FILE_BYTES:
        raise FrameworkRefused(f"the file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 -- any unreadable workbook is a refusal, not a crash
        raise FrameworkRefused(f"not a readable .xlsx workbook ({type(exc).__name__})") from exc
    sheets: dict[str, list[list[Any]]] = {}
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > MAX_ROWS:
                raise FrameworkRefused(f"sheet '{ws.title}' has more than {MAX_ROWS} rows")
            rows.append(list(row))
        sheets[ws.title] = rows
    wb.close()
    return sheets


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _headers(rows: list[list[Any]]) -> list[str]:
    return [_cell(h) for h in (rows[0] if rows else [])]


def propose_mapping(headers: list[str]) -> dict[str, Optional[str]]:
    lower = {h.lower(): h for h in headers if h}
    mapping: dict[str, Optional[str]] = {}
    for field in FIELDS:
        mapping[field] = next((lower[a] for a in _ALIASES[field] if a in lower), None)
    return mapping


def inspect(data: bytes) -> dict:
    """Sheets, their headers and a proposed mapping -- the first step of an
    import, before anything is stored."""
    sheets = _read_workbook(data)
    out = []
    for name, rows in sheets.items():
        headers = _headers(rows)
        out.append({"sheet": name, "headers": headers, "rows": max(len(rows) - 1, 0),
                    "proposed_mapping": propose_mapping(headers)})
    default = next((s["sheet"] for s in out if s["sheet"] == TEMPLATE_SHEET), out[0]["sheet"] if out else None)
    return {"sheets": out, "default_sheet": default, "fields": list(FIELDS), "required_fields": list(REQUIRED_FIELDS),
            "field_labels": FIELD_LABELS}


def parse(data: bytes, sheet: str, mapping: dict[str, Optional[str]], *, derive_parent: bool) -> tuple[list[dict], dict]:
    """Nodes plus a validation report {errors, warnings}. Nothing here
    guesses: a row that cannot be placed is an error, not a best effort."""
    sheets = _read_workbook(data)
    if sheet not in sheets:
        raise FrameworkRefused(f"no sheet named '{sheet}'")
    rows = sheets[sheet]
    headers = _headers(rows)
    errors: list[str] = []
    warnings: list[str] = []
    for field in REQUIRED_FIELDS:
        if not mapping.get(field):
            errors.append(f"map a column to {FIELD_LABELS[field]} (required)")
    for field, col in mapping.items():
        if field not in FIELDS:
            errors.append(f"unknown field '{field}'")
        elif col and col not in headers:
            errors.append(f"column '{col}' (mapped to {FIELD_LABELS[field]}) is not in sheet '{sheet}'")
    if errors:
        return [], {"errors": errors, "warnings": warnings}
    index = {f: headers.index(c) for f, c in mapping.items() if c}

    nodes: list[dict] = []
    seen: dict[str, int] = {}
    for rownum, row in enumerate(rows[1:], start=2):
        values = {f: _cell(row[i]) if i < len(row) else "" for f, i in index.items()}
        if not any(values.values()):
            continue
        key, name = values.get("node_key", ""), values.get("name", "")
        if not key:
            errors.append(f"row {rownum}: Node ID is empty")
            continue
        if not name:
            errors.append(f"row {rownum}: Name is empty for node {key}")
        if key in seen:
            errors.append(f"row {rownum}: Node ID {key} duplicates row {seen[key]}")
            continue
        seen[key] = rownum
        parent = values.get("parent_key", "")
        if not parent and derive_parent and "." in key:
            parent = key.rsplit(".", 1)[0]
        stated_level = values.get("level", "")
        nodes.append({"node_key": key[:100], "parent_key": parent[:100] or None, "name": name[:300],
                      "description": values.get("description", "")[:4000], "node_type": values.get("node_type", "")[:60],
                      "external_ref": values.get("external_ref", "")[:100], "_row": rownum, "_stated_level": stated_level})
    if not nodes and not errors:
        errors.append("the sheet has no node rows")

    by_key = {n["node_key"]: n for n in nodes}
    for n in nodes:
        if n["parent_key"] and n["parent_key"] not in by_key:
            errors.append(f"row {n['_row']}: parent {n['parent_key']} of node {n['node_key']} is not in the file")
            n["parent_key"] = None
    # Depth and cycles.
    for n in nodes:
        depth, cur, path = 1, n, {n["node_key"]}
        while cur["parent_key"]:
            cur = by_key[cur["parent_key"]]
            if cur["node_key"] in path:
                errors.append(f"row {n['_row']}: node {n['node_key']} is part of a parent cycle")
                depth = 0
                break
            path.add(cur["node_key"])
            depth += 1
        n["level"] = depth
        stated = n.pop("_stated_level")
        if stated and depth and stated != str(depth):
            warnings.append(f"row {n['_row']}: Level {stated} differs from its position in the hierarchy ({depth}); "
                            "the hierarchy is used")
    if nodes and not any(n["external_ref"] for n in nodes) and mapping.get("external_ref"):
        warnings.append("the External reference column is mapped but empty")
    if any(n["external_ref"] for n in nodes):
        warnings.append("External references are shown exactly as supplied; Jade does not verify them against APQC "
                        "or any other source")
    order = {n["node_key"]: i for i, n in enumerate(nodes)}
    for n in nodes:
        n["position"] = order[n["node_key"]]
        n.pop("_row", None)
        n["node_sha256"] = node_sha(n)
    return nodes, {"errors": errors, "warnings": warnings}


def _content_sha(nodes: list[dict]) -> str:
    return hashlib.sha256(json.dumps(sorted(n["node_sha256"] for n in nodes)).encode()).hexdigest()


# ---------------------------------------------------------------------
# Frameworks and versions
# ---------------------------------------------------------------------
def _framework_row(r) -> dict:
    return dict(r)


def _version_row(r) -> dict:
    d = dict(r)
    for k in ("column_mapping", "validation", "changes"):
        d[k] = json.loads(d[k] or "{}")
    d["source_statement"] = d["source_statement"]
    return d


def list_frameworks(company_id: str) -> list[dict]:
    with connection() as conn:
        fws = [dict(r) for r in conn.execute(
            "SELECT * FROM process_frameworks WHERE company_id = ? ORDER BY created_at", (company_id,)).fetchall()]
        for f in fws:
            f["versions"] = [_version_row(r) for r in conn.execute(
                "SELECT * FROM framework_versions WHERE framework_id = ? ORDER BY version DESC",
                (f["framework_id"],)).fetchall()]
            f["active_version"] = next((v["version"] for v in f["versions"] if v["status"] == "active"), None)
            f["source_label"] = SOURCE_KINDS.get(f["source_kind"], f["source_kind"])
    return fws


def get_framework(company_id: str, framework_id: str) -> Optional[dict]:
    return next((f for f in list_frameworks(company_id) if f["framework_id"] == framework_id), None)


def get_version(company_id: str, framework_id: str, version: int) -> Optional[dict]:
    with connection() as conn:
        r = conn.execute("SELECT * FROM framework_versions WHERE company_id = ? AND framework_id = ? AND version = ?",
                         (company_id, framework_id, version)).fetchone()
    return _version_row(r) if r else None


def nodes(framework_id: str, version: int) -> list[dict]:
    with connection() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM framework_nodes WHERE framework_id = ? AND version = ? ORDER BY position",
            (framework_id, version)).fetchall()]


def node(framework_id: str, version: int, node_key: str) -> Optional[dict]:
    with connection() as conn:
        r = conn.execute("SELECT * FROM framework_nodes WHERE framework_id = ? AND version = ? AND node_key = ?",
                         (framework_id, version, node_key)).fetchone()
    return dict(r) if r else None


def path_of(framework_id: str, version: int, node_key: str) -> list[dict]:
    """The node and its ancestors, root first."""
    chain, cur, guard = [], node(framework_id, version, node_key), 0
    while cur and guard < 50:
        chain.append({"node_key": cur["node_key"], "name": cur["name"]})
        cur = node(framework_id, version, cur["parent_key"]) if cur["parent_key"] else None
        guard += 1
    return list(reversed(chain))


def _diff(old: list[dict], new: list[dict]) -> dict:
    o = {n["node_key"]: n for n in old}
    n_ = {n["node_key"]: n for n in new}
    return {"added": sorted(set(n_) - set(o)), "removed": sorted(set(o) - set(n_)),
            "changed": sorted(k for k in set(o) & set(n_) if o[k]["node_sha256"] != n_[k]["node_sha256"])}


def active_version(company_id: str, framework_id: str) -> Optional[dict]:
    with connection() as conn:
        r = conn.execute("SELECT * FROM framework_versions WHERE company_id = ? AND framework_id = ? AND status = 'active'",
                         (company_id, framework_id)).fetchone()
    return _version_row(r) if r else None


def create_draft(company_id: str, *, name: str, source_kind: str, source_statement: str, file_name: str, data: bytes,
                 sheet: str, mapping: dict, derive_parent: bool, actor: str,
                 framework_id: Optional[str] = None) -> dict:
    """Store the original file and a draft version (replacing an earlier
    draft of the same framework). A draft is never referenced by anything;
    it becomes an immutable version only on activation."""
    if source_kind not in SOURCE_KINDS:
        raise FrameworkRefused(f"source kind must be one of {', '.join(SOURCE_KINDS)}")
    if not file_name.lower().endswith(".xlsx"):
        raise FrameworkRefused("upload an .xlsx workbook")
    name = name.strip()
    if not name:
        raise FrameworkRefused("give the framework a name")
    if source_kind == "apqc_authorised" and len(source_statement.strip()) < 10:
        raise FrameworkRefused("APQC content needs a statement of the organisation's authority to use it "
                               "(licence or agreement reference)")
    parsed, report = parse(data, sheet, mapping, derive_parent=derive_parent)
    if source_kind == "synthetic_fixture":
        report["warnings"].insert(0, "SYNTHETIC fixture: for demonstration only; it is not APQC content")
    key = artifacts.default_store().put(company_id, data)
    file_sha = hashlib.sha256(data).hexdigest()
    now = _now()
    with connection(immediate=True) as conn:
        if framework_id:
            fw = conn.execute("SELECT * FROM process_frameworks WHERE company_id = ? AND framework_id = ?",
                              (company_id, framework_id)).fetchone()
            if fw is None:
                raise LookupError(f"no such framework: {framework_id}")
            if fw["source_kind"] != source_kind:
                raise FrameworkRefused("a new version must keep the framework's source kind; import a separate "
                                       "framework for different content")
        else:
            if conn.execute("SELECT 1 FROM process_frameworks WHERE company_id = ? AND name = ?",
                            (company_id, name)).fetchone():
                raise FrameworkRefused(f"a framework named '{name}' already exists; import a new version of it instead")
            framework_id = f"PF-{uuid.uuid4().hex[:8]}"
            conn.execute("INSERT INTO process_frameworks (framework_id, company_id, name, source_kind, created_by, "
                         "created_at) VALUES (?, ?, ?, ?, ?, ?)", (framework_id, company_id, name, source_kind, actor, now))
        draft = conn.execute("SELECT version FROM framework_versions WHERE framework_id = ? AND status = 'draft'",
                             (framework_id,)).fetchone()
        if draft:
            version = draft["version"]
            conn.execute("DELETE FROM framework_nodes WHERE framework_id = ? AND version = ?", (framework_id, version))
            conn.execute("DELETE FROM framework_versions WHERE framework_id = ? AND version = ?", (framework_id, version))
        else:
            version = (conn.execute("SELECT MAX(version) AS v FROM framework_versions WHERE framework_id = ?",
                                    (framework_id,)).fetchone()["v"] or 0) + 1
        active = conn.execute("SELECT version FROM framework_versions WHERE framework_id = ? AND status = 'active'",
                              (framework_id,)).fetchone()
        changes = {}
        if active:
            old = [dict(r) for r in conn.execute("SELECT * FROM framework_nodes WHERE framework_id = ? AND version = ?",
                                                 (framework_id, active["version"])).fetchall()]
            changes = {"compared_with_version": active["version"], **_diff(old, parsed)}
        conn.execute(
            "INSERT INTO framework_versions (framework_id, version, company_id, status, file_name, file_sha256, "
            "file_size, storage_key, sheet_name, column_mapping, source_statement, validation, node_count, "
            "content_sha256, changes, uploaded_by, uploaded_at) VALUES (?, ?, ?, 'draft', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (framework_id, version, company_id, file_name[:200], file_sha, len(data), key, sheet,
             json.dumps({"mapping": mapping, "derive_parent": derive_parent}), source_statement.strip()[:1000],
             json.dumps(report), len(parsed), _content_sha(parsed), json.dumps(changes), actor, now))
        conn.executemany(
            "INSERT INTO framework_nodes (framework_id, version, node_key, parent_key, level, position, name, "
            "description, node_type, external_ref, node_sha256) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            [(framework_id, version, n["node_key"], n["parent_key"], n["level"], n["position"], n["name"],
              n["description"], n["node_type"], n["external_ref"], n["node_sha256"]) for n in parsed])
    return preview(company_id, framework_id, version)


def preview(company_id: str, framework_id: str, version: int) -> dict:
    v = get_version(company_id, framework_id, version)
    if v is None:
        raise LookupError(f"no such framework version: {framework_id} v{version}")
    fw = get_framework(company_id, framework_id)
    return {"framework": {k: fw[k] for k in ("framework_id", "name", "source_kind", "source_label", "active_version")},
            "version": v, "nodes": nodes(framework_id, version)}


def remap(company_id: str, framework_id: str, *, sheet: str, mapping: dict, derive_parent: bool, actor: str) -> dict:
    with connection() as conn:
        draft = conn.execute("SELECT * FROM framework_versions WHERE company_id = ? AND framework_id = ? AND status = 'draft'",
                             (company_id, framework_id)).fetchone()
        fw = conn.execute("SELECT * FROM process_frameworks WHERE framework_id = ?", (framework_id,)).fetchone()
    if draft is None:
        raise FrameworkRefused("there is no draft to re-map; upload the file again")
    data = artifacts.default_store().get(draft["storage_key"])
    return create_draft(company_id, name=fw["name"], source_kind=fw["source_kind"],
                        source_statement=draft["source_statement"], file_name=draft["file_name"], data=data,
                        sheet=sheet, mapping=mapping, derive_parent=derive_parent, actor=actor,
                        framework_id=framework_id)


def original_file(company_id: str, framework_id: str, version: int) -> tuple[str, bytes]:
    v = get_version(company_id, framework_id, version)
    if v is None:
        raise LookupError(f"no such framework version: {framework_id} v{version}")
    data = artifacts.default_store().get(v["storage_key"])
    if hashlib.sha256(data).hexdigest() != v["file_sha256"]:
        raise FrameworkRefused("the stored original file no longer matches its recorded checksum")
    return v["file_name"], data


def activate(company_id: str, framework_id: str, version: int, *, actor: str) -> dict:
    """Make a validated draft the active version. The previous version is
    superseded -- kept, with its nodes, for every reference to it."""
    with connection(immediate=True) as conn:
        v = conn.execute("SELECT * FROM framework_versions WHERE company_id = ? AND framework_id = ? AND version = ?",
                         (company_id, framework_id, version)).fetchone()
        if v is None:
            raise LookupError(f"no such framework version: {framework_id} v{version}")
        if v["status"] != "draft":
            raise FrameworkRefused(f"version {version} is {v['status']}; only a draft can be activated")
        report = json.loads(v["validation"])
        if report.get("errors"):
            raise FrameworkRefused(f"version {version} has {len(report['errors'])} validation error(s); fix the file "
                                   "or the column mapping first")
        now = _now()
        conn.execute("UPDATE framework_versions SET status = 'superseded', superseded_at = ? WHERE framework_id = ? "
                     "AND status = 'active'", (now, framework_id))
        conn.execute("UPDATE framework_versions SET status = 'active', activated_by = ?, activated_at = ? "
                     "WHERE framework_id = ? AND version = ?", (actor, now, framework_id, version))
        setting = conn.execute("SELECT * FROM process_settings WHERE company_id = ?", (company_id,)).fetchone()
        if setting is None:
            conn.execute("INSERT INTO process_settings (company_id, selected_framework_id, revision, updated_by, "
                         "updated_at) VALUES (?, ?, 1, ?, ?)", (company_id, framework_id, actor, now))
    changes = json.loads(v["changes"] or "{}")
    from . import story as story_process

    affected = story_process.on_framework_activated(company_id, framework_id, version, changes)
    return {**preview(company_id, framework_id, version), "affected_stories": affected}


# ---------------------------------------------------------------------
# Company setting: the selected framework
# ---------------------------------------------------------------------
def settings(company_id: str) -> dict:
    with connection() as conn:
        r = conn.execute("SELECT * FROM process_settings WHERE company_id = ?", (company_id,)).fetchone()
    return dict(r) if r else {"company_id": company_id, "selected_framework_id": None, "revision": 0}


def select(company_id: str, framework_id: str, *, expected_revision: int, actor: str) -> dict:
    with connection(immediate=True) as conn:
        if not conn.execute("SELECT 1 FROM framework_versions WHERE company_id = ? AND framework_id = ? AND status = 'active'",
                            (company_id, framework_id)).fetchone():
            raise FrameworkRefused("only a framework with an active version can be selected")
        r = conn.execute("SELECT revision FROM process_settings WHERE company_id = ?", (company_id,)).fetchone()
        current = r["revision"] if r else 0
        if current != expected_revision:
            raise FrameworkRefused(f"the setting changed meanwhile (revision {current}); reload and try again")
        if r:
            conn.execute("UPDATE process_settings SET selected_framework_id = ?, revision = ?, updated_by = ?, "
                         "updated_at = ? WHERE company_id = ?", (framework_id, current + 1, actor, _now(), company_id))
        else:
            conn.execute("INSERT INTO process_settings (company_id, selected_framework_id, revision, updated_by, "
                         "updated_at) VALUES (?, ?, 1, ?, ?)", (company_id, framework_id, actor, _now()))
    return settings(company_id)


def selected_active(company_id: str) -> Optional[dict]:
    """The selected framework and its active version, or None."""
    fid = settings(company_id).get("selected_framework_id")
    if not fid:
        return None
    v = active_version(company_id, fid)
    fw = get_framework(company_id, fid)
    return {"framework": fw, "version": v} if v and fw else None


def search(framework_id: str, version: int, query: str, limit: int = 25) -> list[dict]:
    words = [w for w in re.split(r"\W+", query.lower()) if len(w) > 2]
    scored = []
    for n in nodes(framework_id, version):
        hay = f"{n['node_key']} {n['name']} {n['description']}".lower()
        score = sum(hay.count(w) for w in words)
        if score:
            scored.append((score, n))
    scored.sort(key=lambda t: (-t[0], t[1]["position"]))
    return [n for _, n in scored[:limit]]
