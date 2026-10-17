"""
A process framework workbook laid out like an APQC PCF export (PCF ID,
Hierarchy ID, Name; categories numbered "4.0"; a separate glossary sheet with
a Definition per Hierarchy ID) imports without manual mapping, and the
Administrator turns chosen processes of the active version into business
domains. The workbook here is SYNTHETIC -- invented names, not APQC content.
"""

from __future__ import annotations

import base64
import io

from .conftest import headers

H = headers("vdb")


def _workbook() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Introduction"
    wb.active.append(["SYNTHETIC test workbook"])
    rows = [(90001, "1.0", "SYN Plan the business"), (90002, "1.1", "SYN Set direction"),
            (90003, "1.1.1", "SYN Scan the market"), (90004, "2.0", "SYN Make and ship goods"),
            (90005, "2.1", "SYN Plan supply"), (90006, "2.2", "SYN Ship orders"), (90007, "2.2.1", "SYN Pick and pack")]
    for title in ("1.0", "Combined"):
        ws = wb.create_sheet(title)
        ws.append(["PCF ID", "Hierarchy ID", "Name", "Difference Index", "Change details", "Metrics available?"])
        for pcf, hid, name in (rows if title == "Combined" else rows[:3]):
            ws.append([pcf, hid, name, 0, None, "N"])
    g = wb.create_sheet("Glossary terms")
    g.append(["Process Element ID", "Hierarchy ID", "Process element", "Definition"])
    for pcf, hid, name in rows:
        g.append([pcf, hid, name, f"Definition of {name}"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(client) -> dict:
    data = _workbook()
    b64 = base64.b64encode(data).decode()
    i = client.post("/process/frameworks/inspect", headers=H, json={"fileName": "pcf.xlsx", "contentBase64": b64}).json()
    assert i["default_sheet"] == "Combined" and i["looks_like_apqc"] and i["description_source"] == "Glossary terms"
    sheet = next(s for s in i["sheets"] if s["sheet"] == "Combined")
    assert sheet["proposed_mapping"]["node_key"] == "Hierarchy ID" and sheet["proposed_mapping"]["external_ref"] == "PCF ID"
    assert sheet["proposed_derive_parent"] is True
    r = client.post("/process/frameworks/drafts", headers=H, json={
        "name": "SYN PCF", "sourceKind": "customer_defined", "sourceStatement": "test", "fileName": "pcf.xlsx",
        "contentBase64": b64, "sheet": "Combined", "mapping": sheet["proposed_mapping"], "deriveParent": True})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["version"]["validation"]["errors"] == [], d["version"]["validation"]
    return d


def test_an_apqc_style_workbook_imports_with_its_numbering_and_glossary(client):
    d = _import(client)
    nodes = {n["node_key"]: n for n in d["nodes"]}
    assert nodes["1.0"]["parent_key"] is None and nodes["2.0"]["parent_key"] is None
    assert nodes["1.1"]["parent_key"] == "1.0" and nodes["2.2.1"]["parent_key"] == "2.2"
    assert (nodes["1.0"]["level"], nodes["2.2"]["level"], nodes["2.2.1"]["level"]) == (1, 2, 3)
    assert nodes["2.2"]["description"] == "Definition of SYN Ship orders" and nodes["2.2"]["external_ref"] == "90006"
    assert any("Glossary terms" in w for w in d["version"]["validation"]["warnings"])


def test_business_domains_are_created_from_an_active_version_without_duplicates(client):
    d = _import(client)
    fid, v = d["framework"]["framework_id"], d["version"]["version"]
    url = f"/process/frameworks/{fid}/versions/{v}/domains"
    # Only from an activated version.
    r = client.post(url, headers=H, json={"nodeKeys": ["1.0"]})
    assert r.status_code == 422 and "activate" in r.json()["detail"]
    assert client.post(f"/process/frameworks/{fid}/versions/{v}/activate", headers=H).status_code == 200
    r = client.post(url, headers=H, json={"nodeKeys": ["2.1", "2.2"]})
    assert r.status_code == 200, r.text
    assert [c["code"] for c in r.json()["created"]] == ["2.1", "2.2"]
    domains = {x["apqcCode"]: x for x in client.get("/business-domains", headers=H).json()}
    assert domains["2.2"]["name"] == "SYN Ship orders" and domains["2.2"]["level"] == "2.2"
    assert domains["2.2"]["description"] == "Definition of SYN Ship orders" and domains["2.2"]["status"] == "active"
    # Running it again never duplicates or overwrites.
    r = client.post(url, headers=H, json={"nodeKeys": ["2.2", "1.1"]}).json()
    assert [c["code"] for c in r["created"]] == ["1.1"] and [c["code"] for c in r["existing"]] == ["2.2"]
    assert sum(1 for x in client.get("/business-domains", headers=H).json() if x["apqcCode"] == "2.2") == 1
    # Unknown processes are refused; other companies see nothing.
    assert client.post(url, headers=H, json={"nodeKeys": ["9.9"]}).status_code == 422
    assert all(x["apqcCode"] not in ("2.1", "2.2") for x in client.get("/business-domains", headers=headers("nhd")).json())
